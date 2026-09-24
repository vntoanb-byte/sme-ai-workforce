"""Kiểm thử app/tools — từng công cụ chạy độc lập trên DB/kho tệp tạm."""

from __future__ import annotations

import io
import os
import time
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.adapters.storage_local import LocalFileStorage
from app.core.errors import ValidationFailed
from app.models.artifact import Artifact, Document
from app.models.audit import AuditLog, LlmCall
from app.models.extraction import Extraction, QCResult
from app.models.workflow import Tool
from app.tools.base import (
    ToolContext,
    artifact_ref,
    get_tool,
    load_all_tools,
    resolve_allowed_path,
    sync_tools_to_db,
    tool_catalog,
)
from app.tools.xlsx_tools import normalize_cell
from tests.conftest import SmartLLM, invoice_payload, png_bytes


def _ctx(db: Session, storage: LocalFileStorage, llm: object = None, **kw) -> ToolContext:
    logs: list[tuple[str, str]] = []
    ctx = ToolContext(db=db, storage=storage, llm=llm or SmartLLM(), run_id=None, **kw)  # type: ignore[arg-type]
    ctx.log = lambda level, msg: logs.append((level, msg))
    ctx.logs = logs  # type: ignore[attr-defined]
    return ctx


def _xlsx(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_artifact(db: Session, storage: LocalFileStorage, artifact_id: int) -> list[list[object]]:
    artifact = db.get(Artifact, artifact_id)
    assert artifact is not None
    with storage.open(artifact_ref(artifact)) as fh:
        wb = load_workbook(io.BytesIO(fh.read()))
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


# ─── registry ───


def test_registry_and_db_sync(db: Session) -> None:
    registry = load_all_tools()
    assert len(registry) == 14
    rows = {t.code: t for t in db.scalars(select(Tool))}
    assert set(rows) == set(registry)

    # Quản trị viên tắt công cụ → validators thấy is_enabled=False; đồng bộ lại
    # không tự bật lại. Mã không còn trong registry bị tắt, không xoá.
    rows["xlsx.dedupe"].is_enabled = False
    db.add(Tool(code="cu.da_bo", name="Cũ", is_enabled=True))
    db.flush()
    sync_tools_to_db(db)
    catalog = tool_catalog(db)
    assert catalog["xlsx.dedupe"]["is_enabled"] is False
    assert db.scalar(select(Tool).where(Tool.code == "cu.da_bo")).is_enabled is False  # type: ignore[union-attr]
    assert catalog["fs.list_new_files"]["required_params"] == ["path"]


def test_path_guard(scan_dir: Path, tmp_path: Path) -> None:
    assert resolve_allowed_path("sub/a.xlsx") == (scan_dir / "sub" / "a.xlsx").resolve()
    assert resolve_allowed_path(str(scan_dir)) == scan_dir.resolve()
    for bad in ("/etc/passwd", str(tmp_path / "khac"), "../ngoai.xlsx"):
        with pytest.raises(ValidationFailed):
            resolve_allowed_path(bad)


# ─── fs ───


def test_list_new_files_filters_extension_and_time(
    db: Session, storage: LocalFileStorage, scan_dir: Path
) -> None:
    old = scan_dir / "cu.jpg"
    old.write_bytes(b"x")
    past = time.time() - 3600
    os.utime(old, (past, past))
    (scan_dir / "moi.pdf").write_bytes(b"%PDF")
    (scan_dir / "ghichu.txt").write_bytes(b"t")
    (scan_dir / "rong.png").write_bytes(b"")  # 0 byte = đang ghi dở
    (scan_dir / ".an.png").write_bytes(b"x")

    tool = get_tool("fs.list_new_files")
    ctx = _ctx(db, storage, since=datetime.now(UTC) - timedelta(minutes=10))
    out = tool.run(ctx, {"path": str(scan_dir), "stable_seconds": 0}, {})
    assert [f["filename"] for f in out.outputs["files"]] == ["moi.pdf"]

    everything = tool.run(
        _ctx(db, storage), {"path": str(scan_dir), "stable_seconds": 0, "only_new": False}, {}
    )
    assert [f["filename"] for f in everything.outputs["files"]] == ["cu.jpg", "moi.pdf"]


def test_list_new_files_rejects_outside_folder(db: Session, storage: LocalFileStorage) -> None:
    with pytest.raises(ValidationFailed):
        get_tool("fs.list_new_files").run(_ctx(db, storage), {"path": "/etc"}, {})


# ─── vision + qc ───


def test_extract_invoice_then_validate(db: Session, storage: LocalFileStorage, scan_dir: Path) -> None:
    for i in range(3):
        (scan_dir / f"hd{i}.png").write_bytes(png_bytes(i))
    files = [{"filename": f"hd{i}.png", "path": str(scan_dir / f"hd{i}.png")} for i in range(3)]
    llm = SmartLLM(
        invoices=[invoice_payload("1"), invoice_payload("2", bad_total=True), invoice_payload("3")]
    )
    ctx = _ctx(db, storage, llm)
    out = get_tool("vision.extract_invoice").run(ctx, {}, {"files": files})
    ids = out.outputs["document_ids"]
    assert len(ids) == 3
    assert all(db.get(Document, i).status == "processing" for i in ids)  # type: ignore[union-attr]
    assert db.scalar(select(LlmCall).limit(1)) is not None

    qc = get_tool("qc.validate_invoice").run(ctx, {}, {"document_ids": ids})
    assert qc.outputs["passed_document_ids"] == [ids[0], ids[2]]
    assert qc.outputs["review_document_ids"] == [ids[1]]
    assert db.get(Document, ids[1]).status == "needs_review"  # type: ignore[union-attr]
    assert qc.detail == "Đạt 2 · cần xác nhận 1"

    # Chạy lại QC không nhân đôi kết quả (bất biến khi lặp).
    get_tool("qc.validate_invoice").run(ctx, {}, {"document_ids": ids})
    extraction = db.scalar(select(Extraction).where(Extraction.document_id == ids[0]))
    count = len(db.scalars(select(QCResult).where(QCResult.extraction_id == extraction.id)).all())  # type: ignore[union-attr]
    assert count == 8

    # Tệp đã xử lý xong được bỏ qua ở lần quét sau (khử trùng sha256).
    again = get_tool("vision.extract_invoice").run(ctx, {}, {"files": files})
    assert again.outputs["document_ids"] == []
    assert "bỏ qua 3 tệp trùng" in (again.detail or "")


def test_escalate_marks_needs_review_and_audits(db: Session, storage: LocalFileStorage) -> None:
    artifact = Artifact(sha256="a" * 64, path="x", size_bytes=1)
    doc = Document(filename="a.png", source_kind="image", status="processing", artifact=artifact)
    db.add(doc)
    db.flush()
    out = get_tool("qc.escalate").run(_ctx(db, storage), {"assign_to": "quanly"}, {"document_ids": [doc.id]})
    assert doc.status == "needs_review"
    assert out.outputs["review_document_ids"] == [doc.id]
    assert db.scalar(select(AuditLog).where(AuditLog.action == "document.escalate")) is not None


def test_classify_document(db: Session, storage: LocalFileStorage, scan_dir: Path) -> None:
    (scan_dir / "a.png").write_bytes(png_bytes(1))
    (scan_dir / "b.txt").write_bytes(b"x")
    files = [
        {"filename": "a.png", "path": str(scan_dir / "a.png")},
        {"filename": "b.txt", "path": str(scan_dir / "b.txt")},
    ]
    out = get_tool("vision.classify_document").run(
        _ctx(db, storage, SmartLLM(label="hop_dong")), {"labels": "hoa_don, hop_dong"},
        {"files": files},
    )
    assert out.outputs["classifications"] == [
        {"filename": "a.png", "label": "hop_dong"},
        {"filename": "b.txt", "label": "khac"},
    ]
    rows = _read_artifact(db, storage, out.outputs["artifact_ids"][0])
    assert rows[0] == ["Tệp", "Loại tài liệu"]


# ─── xlsx ───


def test_append_rows_never_overwrites_target(db: Session, storage: LocalFileStorage, scan_dir: Path) -> None:
    target = scan_dir / "SoHoaDon.xlsx"
    original = _xlsx([["Số hoá đơn", "Ghi chú"], ["0000001", "có sẵn"]])
    target.write_bytes(original)

    (scan_dir / "hd.png").write_bytes(png_bytes(5))
    ctx = _ctx(db, storage, SmartLLM(invoices=[invoice_payload("0000777")]))
    ids = get_tool("vision.extract_invoice").run(
        ctx, {}, {"files": [{"filename": "hd.png", "path": str(scan_dir / "hd.png")}]}
    ).outputs["document_ids"]

    out = get_tool("xlsx.append_rows").run(ctx, {"file": "SoHoaDon.xlsx", "sheet": "Sheet"},
                                           {"document_ids": ids})
    assert target.read_bytes() == original  # tệp gốc KHÔNG bị đụng tới
    assert out.outputs["workbook"]["filename"].startswith("SoHoaDon_")
    rows = _read_artifact(db, storage, out.outputs["artifact_ids"][0])
    assert rows[1][0] == "0000001"
    assert rows[2][0] == "0000777" and rows[2][-1] == "hd.png"
    assert Decimal(str(rows[2][8])) == Decimal("110000")
    assert db.scalar(select(AuditLog).where(AuditLog.action == "run.output")) is not None

    # Tệp đích chưa tồn tại → bảng mới, tiêu đề đúng dòng 1, dữ liệu từ dòng 2.
    fresh = get_tool("xlsx.append_rows").run(ctx, {"file": "Moi.xlsx"}, {"document_ids": ids})
    rows = _read_artifact(db, storage, fresh.outputs["artifact_ids"][0])
    assert rows[0][0] == "Số hoá đơn" and rows[1][0] == "0000777" and len(rows) == 2


def test_merge_normalize_dedupe_chain(db: Session, storage: LocalFileStorage, scan_dir: Path) -> None:
    (scan_dir / "a.xlsx").write_bytes(_xlsx([["Mã", "Tên", "Tiền"], ["K1", "  An   Phát ", "1.200.000"]]))
    (scan_dir / "b.xlsx").write_bytes(
        _xlsx([["Mã", "Ngày", "Tiền"], ["K1", "05/08/2026", "1.200.000"], ["K2", None, "500"]])
    )
    files = [{"filename": n, "path": str(scan_dir / n)} for n in ("a.xlsx", "b.xlsx")]
    ctx = _ctx(db, storage)
    merged = get_tool("xlsx.merge_files").run(ctx, {}, {"files": files}).outputs
    assert _read_artifact(db, storage, merged["workbook"]["artifact_id"])[0] == ["Mã", "Tên", "Tiền", "Ngày"]

    normalized = get_tool("xlsx.normalize").run(ctx, {}, merged).outputs
    rows = _read_artifact(db, storage, normalized["workbook"]["artifact_id"])
    assert rows[1][1] == "An Phát"
    assert rows[1][2] == 1200000
    assert rows[2][3].date().isoformat() == "2026-08-05"  # type: ignore[union-attr]

    deduped = get_tool("xlsx.dedupe").run(ctx, {"key_columns": "Mã"}, normalized)
    rows = _read_artifact(db, storage, deduped.outputs["workbook"]["artifact_id"])
    assert [r[0] for r in rows[1:]] == ["K1", "K2"]
    assert deduped.detail == "Bỏ 1 dòng trùng, còn 2 dòng"
    assert len(deduped.outputs["artifact_ids"]) == 3

    with pytest.raises(ValidationFailed):
        get_tool("xlsx.dedupe").run(ctx, {"key_columns": "KhongCo"}, normalized)


def test_normalize_cell() -> None:
    assert normalize_cell("  a   b ") == "a b"
    assert normalize_cell("1.234.567 đ") == Decimal("1234567")
    assert normalize_cell("Ngày 01 tháng 02 năm 2026").isoformat() == "2026-02-01"  # type: ignore[union-attr]
    assert normalize_cell("HD-001") == "HD-001"
    assert normalize_cell("   ") is None
    assert normalize_cell(5) == 5


def test_reconcile(db: Session, storage: LocalFileStorage, scan_dir: Path) -> None:
    (scan_dir / "so.xlsx").write_bytes(
        _xlsx([["Số HĐ", "Tiền"], ["001", 100000], ["002", 200000], ["003", 50]])
    )
    (scan_dir / "congno.xlsx").write_bytes(
        _xlsx([["Số HĐ", "Tiền"], ["001", "100.000"], ["002", 210000], ["004", 1]])
    )
    out = get_tool("xlsx.reconcile").run(
        _ctx(db, storage),
        {"left_file": "so.xlsx", "right_file": "congno.xlsx", "key_column": "Số HĐ"},
        {},
    )
    assert out.outputs["mismatch_count"] == 3
    assert out.detail == "1 ô lệch, 1 dòng chỉ có bên trái, 1 dòng chỉ có bên phải"
    rows = _read_artifact(db, storage, out.outputs["artifact_ids"][0])
    assert rows[1][:2] == ["002", "Tiền"]

    with pytest.raises(ValidationFailed):
        get_tool("xlsx.reconcile").run(
            _ctx(db, storage),
            {"left_file": "so.xlsx", "right_file": "congno.xlsx", "key_column": "Mã"},
            {},
        )


# ─── report + doc ───


def test_report_tools_create_outputs(db: Session, storage: LocalFileStorage) -> None:
    ctx = _ctx(db, storage)
    xlsx = get_tool("report.build_xlsx").run(ctx, {"period": "this_month"}, {})
    pdf = get_tool("report.build_pdf").run(ctx, {"period": "this_month"}, xlsx.outputs)
    assert len(pdf.outputs["artifact_ids"]) == 2
    artifact = db.get(Artifact, pdf.outputs["artifact_ids"][1])
    with storage.open(artifact_ref(artifact)) as fh:  # type: ignore[arg-type]
        assert fh.read(4) == b"%PDF"


def test_doc_tools(db: Session, storage: LocalFileStorage, scan_dir: Path) -> None:
    from PIL import Image

    img = scan_dir / "to.png"
    Image.new("RGB", (3000, 2000), "white").save(img)
    pdf = scan_dir / "hai_trang.pdf"
    Image.new("RGB", (200, 300), "white").save(
        pdf, save_all=True, append_images=[Image.new("RGB", (200, 300), "gray")]
    )
    ctx = _ctx(db, storage)
    files = [{"filename": "to.png", "path": str(img)}, {"filename": "hai_trang.pdf", "path": str(pdf)}]
    rendered = get_tool("doc.render_pdf").run(ctx, {}, {"files": files})
    names = [f["filename"] for f in rendered.outputs["files"]]
    assert names == ["to.png", "hai_trang_p1.png", "hai_trang_p2.png"]

    processed = get_tool("doc.preprocess_image").run(ctx, {}, rendered.outputs)
    first = db.get(Artifact, processed.outputs["files"][0]["artifact_id"])
    with storage.open(artifact_ref(first)) as fh:  # type: ignore[arg-type]
        out = Image.open(io.BytesIO(fh.read()))
    assert max(out.size) <= 1280
