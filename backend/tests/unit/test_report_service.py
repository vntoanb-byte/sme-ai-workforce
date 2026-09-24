"""Kiểm thử services/report_service.py — tổng hợp SQL, kỳ báo cáo, kết xuất."""

from __future__ import annotations

import io
import json
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.errors import ValidationFailed
from app.models.artifact import Artifact, Document
from app.models.extraction import Extraction
from app.services import report_service


def _doc(
    db: Session, n: int, *, seller: str, day: date, rate: str, subtotal: str, status: str = "ok"
) -> Document:
    artifact = Artifact(sha256=f"{n:064d}", path=f"p{n}", size_bytes=1)
    doc = Document(filename=f"{n}.png", source_kind="image", status=status, artifact=artifact)
    db.add(doc)
    db.flush()
    sub = Decimal(subtotal)
    vat = sub * Decimal(rate) / 100
    db.add(
        Extraction(
            document_id=doc.id, schema_version="invoice_v1", model_name="m", latency_ms=1,
            invoice_no=str(n), issue_date=day, seller_name=seller, currency="VND",
            subtotal=sub, vat_rate=Decimal(rate), vat_amount=vat, total=sub + vat,
            extracted_data_json=json.dumps({}),
        )
    )
    db.flush()
    return doc


@pytest.fixture()
def seeded(db: Session) -> Session:
    _doc(db, 1, seller="Minh Long", day=date(2026, 8, 3), rate="10", subtotal="1000000")
    _doc(db, 2, seller="Minh Long", day=date(2026, 8, 20), rate="8", subtotal="500000")
    _doc(db, 3, seller="Đông Á", day=date(2026, 9, 1), rate="10", subtotal="200000")
    _doc(db, 4, seller="Đông Á", day=date(2026, 8, 5), rate="10", subtotal="999999",
         status="needs_review")  # chưa xác nhận → không tính
    _doc(db, 5, seller="Ngoài kỳ", day=date(2026, 6, 1), rate="10", subtotal="1")
    # Chứng từ 2 được sửa tay: chỉ extraction MỚI NHẤT được tính.
    doc2 = db.get(Document, 2)
    db.add(
        Extraction(
            document_id=doc2.id, schema_version="invoice_v1", model_name="human", latency_ms=0,  # type: ignore[union-attr]
            invoice_no="2", issue_date=date(2026, 8, 20), seller_name="Minh Long",
            currency="VND", subtotal=Decimal("600000"), vat_rate=Decimal("8"),
            vat_amount=Decimal("48000"), total=Decimal("648000"), extracted_data_json="{}",
        )
    )
    db.flush()
    return db


def test_aggregate_by_seller(seeded: Session) -> None:
    data = report_service.aggregate(seeded, date(2026, 8, 1), date(2026, 9, 30), "seller")
    rows = {r.group: r for r in data.rows}
    assert set(rows) == {"Minh Long", "Đông Á"}
    assert rows["Minh Long"].doc_count == 2
    assert rows["Minh Long"].subtotal == Decimal("1600000.00")
    assert rows["Minh Long"].total == Decimal("1748000.00")
    assert rows["Đông Á"].total == Decimal("220000.00")
    assert data.grand_total == Decimal("1968000.00")
    assert isinstance(data.grand_total, Decimal)


def test_aggregate_by_month_and_vat_rate(seeded: Session) -> None:
    by_month = report_service.aggregate(seeded, date(2026, 8, 1), date(2026, 9, 30), "month")
    assert {r.group: r.doc_count for r in by_month.rows} == {"08/2026": 2, "09/2026": 1}
    by_rate = report_service.aggregate(seeded, date(2026, 8, 1), date(2026, 9, 30), "vat_rate")
    assert {r.group: r.doc_count for r in by_rate.rows} == {"10%": 2, "8%": 1}


def test_aggregate_validates_input(db: Session) -> None:
    with pytest.raises(ValidationFailed):
        report_service.aggregate(db, date(2026, 9, 1), date(2026, 8, 1), "seller")
    with pytest.raises(ValidationFailed):
        report_service.aggregate(db, date(2026, 8, 1), date(2026, 9, 1), "khong_co")


def test_resolve_period() -> None:
    today = date(2026, 9, 24)  # thứ Năm
    assert report_service.resolve_period("today", today) == (today, today)
    assert report_service.resolve_period("this_week", today) == (date(2026, 9, 21), today)
    assert report_service.resolve_period("this_month", today) == (date(2026, 9, 1), today)
    assert report_service.resolve_period("last_month", today) == (date(2026, 8, 1), date(2026, 8, 31))
    with pytest.raises(ValidationFailed):
        report_service.resolve_period("nam_ngoai", today)


def test_build_files(seeded: Session) -> None:
    data = report_service.aggregate(seeded, date(2026, 8, 1), date(2026, 9, 30), "seller")
    wb = load_workbook(io.BytesIO(report_service.build_xlsx(data)))
    rows = [r for r in wb.active.iter_rows(values_only=True)]
    assert rows[2] == ("Nhà cung cấp", "Số chứng từ", "Tiền hàng", "Tiền thuế", "Tổng cộng")
    assert rows[-1][0] == "Tổng cộng" and rows[-1][1] == 3
    assert report_service.build_pdf(data)[:4] == b"%PDF"
