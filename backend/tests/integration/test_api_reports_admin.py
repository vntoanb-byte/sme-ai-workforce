"""Kiểm thử API báo cáo + quản trị (người dùng, cấu hình, chỉ số, thử kết nối mô hình)."""

from __future__ import annotations

import io
from datetime import date, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.deps import get_llm
from app.core.config import settings
from app.main import app
from app.ports.llm import LLMUnavailable
from tests.conftest import PASSWORD, FakeLLM, SmartLLM, invoice_payload, png_bytes


def _upload(client: TestClient, fake_llm: FakeLLM, payload: dict[str, Any], seed: int) -> dict:
    fake_llm._parsed = payload
    return client.post(
        "/api/v1/documents", files={"file": (f"{seed}.png", png_bytes(seed), "image/png")}
    ).json()


def _period() -> dict[str, str]:
    today = date.today()
    return {"from": (today - timedelta(days=30)).isoformat(), "to": today.isoformat()}


def test_report_preview_and_export(client: TestClient, fake_llm: FakeLLM) -> None:
    _upload(client, fake_llm, invoice_payload("0000101"), 1)
    _upload(client, fake_llm, invoice_payload("0000102"), 2)
    _upload(client, fake_llm, invoice_payload("0000103", bad_total=True), 3)  # chưa xác nhận

    preview = client.post("/api/v1/reports/preview", json={**_period(), "group_by": "seller"})
    assert preview.status_code == 200
    body = preview.json()
    assert body["from"] == _period()["from"] and body["group_by"] == "seller"
    assert body["rows"] == [
        {
            "group": "Công ty TNHH ABC",
            "doc_count": 2,
            "subtotal": 200000.0,
            "vat_amount": 20000.0,
            "total": 220000.0,
        }
    ]
    assert body["grand_total"] == 220000.0

    for fmt, magic in (("xlsx", b"PK"), ("pdf", b"%PDF")):
        exported = client.post("/api/v1/reports/export", json={**_period(), "format": fmt})
        assert exported.status_code == 200
        result = exported.json()
        assert result["filename"].endswith(f".{fmt}")
        download = client.get(result["download_url"])
        assert download.status_code == 200
        assert download.content[: len(magic)] == magic
        assert "attachment" in download.headers["content-disposition"]
    rows = list(
        load_workbook(
            io.BytesIO(
                client.get(
                    client.post("/api/v1/reports/export", json=_period()).json()["download_url"]
                ).content
            )
        ).active.iter_rows(values_only=True)
    )
    assert rows[3][0] == "Công ty TNHH ABC"

    bad = client.post("/api/v1/reports/preview", json={"from": "2026-09-01", "to": "2026-08-01"})
    assert bad.status_code == 422
    assert client.get("/api/v1/reports/99999/download").status_code == 404


def test_admin_user_management(client: TestClient, users) -> None:
    listed = client.get("/api/v1/admin/users").json()
    assert [u["username"] for u in listed] == ["ketoan", "quanly", "admin"]

    created = client.post(
        "/api/v1/admin/users",
        json={
            "username": "nhanvien",
            "full_name": "Lê Văn Minh",
            "email": "minh@example.vn",
            "password": "matkhau-moi-1",
            "roles": ["MANAGER"],
        },
    )
    assert created.status_code == 201
    assert created.json()["roles"] == ["USER", "MANAGER"]  # luôn có quyền USER cơ bản
    dup = client.post(
        "/api/v1/admin/users",
        json={
            "username": "NhanVien",
            "full_name": "x",
            "email": "khac@example.vn",
            "password": "matkhau-moi-1",
        },
    )
    assert dup.status_code == 409
    weak = client.post(
        "/api/v1/admin/users",
        json={
            "username": "yeu",
            "full_name": "x",
            "email": "y@example.vn",
            "password": "123",
        },
    )
    assert weak.status_code == 422

    uid = created.json()["id"]
    changed = client.patch(
        f"/api/v1/admin/users/{uid}", json={"roles": ["USER"], "password": "doi-mat-khau-2"}
    )
    assert changed.json()["roles"] == ["USER"]
    login = client.post(
        "/api/v1/auth/login", json={"username": "nhanvien", "password": "doi-mat-khau-2"}
    )
    assert login.status_code == 200

    me = users["admin"].id
    assert client.patch(f"/api/v1/admin/users/{me}", json={"is_active": False}).status_code == 409
    assert client.patch(f"/api/v1/admin/users/{me}", json={"roles": ["USER"]}).status_code == 409
    assert client.patch("/api/v1/admin/users/9999", json={"full_name": "x"}).status_code == 404
    assert (
        client.post(
            "/api/v1/auth/login", json={"username": "admin", "password": PASSWORD}
        ).status_code
        == 200
    )


def test_settings_with_secret_masking(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "CREDENTIAL_ENC_KEY", "khoa-ma-hoa-thu-nghiem")
    initial = {s["key"]: s["value"] for s in client.get("/api/v1/admin/settings").json()}
    assert initial["manual_minutes_per_doc"] == 4

    resp = client.put(
        "/api/v1/admin/settings",
        json={
            "company_name": "Công ty TNHH TM An Phát",
            "secret.smtp_password": "bi-mat-123",
        },
    )
    values = {s["key"]: s["value"] for s in resp.json()}
    assert values["company_name"] == "Công ty TNHH TM An Phát"
    assert values["secret.smtp_password"] == "••••••"
    assert "bi-mat-123" not in resp.text

    bad = client.put("/api/v1/admin/settings", json={"Khoa Sai": 1})
    assert bad.status_code == 422


def test_metrics(client: TestClient, fake_llm: FakeLLM) -> None:
    _upload(client, fake_llm, invoice_payload("0000201"), 1)
    _upload(client, fake_llm, invoice_payload("0000202"), 2)
    bad = _upload(client, fake_llm, invoice_payload("0000203", bad_total=True), 3)
    client.post(f"/api/v1/reviews/{bad['id']}/resolve", json={"action": "approve"})

    m = client.get("/api/v1/admin/metrics").json()
    assert m["docs_processed"] == 3
    assert m["automation_rate"] == pytest.approx(66.7)
    assert m["hours_saved"] == pytest.approx(round(2 * 4 / 60, 1))
    assert m["minutes_per_doc"] == 4
    assert m["needs_review_total"] == 1 and m["needs_review_open"] == 0
    assert m["llm_status"] == "ok"
    assert m["tokens_this_month"] == 60  # 3 lời gọi × (10 + 10) token (FakeLLM)
    assert m["period_label"].startswith("Tháng ")


def test_llm_status_and_connection_test(client: TestClient, fake_llm: FakeLLM) -> None:
    ok = SmartLLM()
    app.dependency_overrides[get_llm] = lambda: ok
    result = client.post("/api/v1/admin/llm/test").json()
    assert result["ok"] is True and result["model"] == "fake-vl"

    down = FakeLLM(raise_exc=LLMUnavailable("Bộ ngắt mạch đang mở"))
    app.dependency_overrides[get_llm] = lambda: down
    result = client.post("/api/v1/admin/llm/test").json()
    assert result["ok"] is False and "ngắt mạch" in result["error"]
    # 1 lời gọi thành công + 1 lỗi trong 15 phút → degraded
    assert client.get("/api/v1/admin/metrics").json()["llm_status"] == "degraded"
