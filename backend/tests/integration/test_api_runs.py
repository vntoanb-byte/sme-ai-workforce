"""Kiểm thử API lần chạy: kích hoạt, danh sách/chi tiết, huỷ, SSE nhật ký, tải tệp."""

from __future__ import annotations

import io
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.workers.worker import Worker
from tests.conftest import SmartLLM, invoice_payload, png_bytes


def _run_worker(queue, session_factory, storage, llm) -> None:  # noqa: ANN001
    assert Worker(
        queue, session_factory, storage, llm, worker_id="w", poll_interval=0, retry_delay=0
    ).process_one()


def test_trigger_list_detail_and_logs(
    client: TestClient, make_employee, scan_dir: Path, queue, session_factory, storage
) -> None:
    employee = make_employee()
    (scan_dir / "hd.png").write_bytes(png_bytes(1))

    created = client.post("/api/v1/runs", json={"employee_id": employee.id})
    assert created.status_code == 201
    run_id = created.json()["id"]
    # Không cho chạy chồng khi lần trước chưa kết thúc.
    again = client.post("/api/v1/runs", json={"employee_id": employee.id})
    assert again.status_code == 409 and again.json()["error"]["code"] == "CONFLICT"

    row = client.get(f"/api/v1/runs/{run_id}").json()
    assert row["status"] == "PENDING" and row["trigger_type"] == "manual"

    _run_worker(queue, session_factory, storage, SmartLLM(invoices=[invoice_payload("0000501")]))

    listing = client.get("/api/v1/runs", params={"status": "SUCCEEDED"}).json()
    assert listing["total"] == 1
    item = listing["items"][0]
    assert item["employee_name"] == "Kế toán hoá đơn" and item["doc_count"] == 1
    assert item["started_at"].endswith("Z") and item["finished_at"].endswith("Z")
    assert client.get("/api/v1/runs", params={"employee_id": employee.id + 99}).json()["total"] == 0
    assert client.get("/api/v1/runs", params={"status": "KHONG_CO"}).status_code == 422

    detail = client.get(f"/api/v1/runs/{run_id}").json()
    assert [s["status"] for s in detail["steps"]] == [
        "SUCCEEDED",
        "SUCCEEDED",
        "SUCCEEDED",
        "SUCCEEDED",
        "SKIPPED",
    ]
    assert detail["stats"] == {"read": 1, "passed": 1, "needs_review": 0, "total_amount": 110000.0}
    output = detail["outputs"][0]

    # Tải tệp Excel kết quả — đúng nội dung, đúng tên tệp.
    download = client.get(output["download_url"], params={"filename": output["filename"]})
    assert download.status_code == 200
    assert output["filename"] in download.headers["content-disposition"].replace("%5F", "_")
    rows = list(load_workbook(io.BytesIO(download.content)).active.iter_rows(values_only=True))
    assert rows[1][0] == "0000501"

    # SSE: phát lại toàn bộ nhật ký rồi báo 'done' vì lần chạy đã kết thúc.
    with client.stream("GET", f"/api/v1/runs/{run_id}/logs") as stream:
        assert stream.headers["content-type"].startswith("text/event-stream")
        payload = "".join(stream.iter_text())
    events = [block for block in payload.split("\n\n") if block.startswith("id:")]
    lines = [json.loads(e.split("data: ", 1)[1]) for e in events]
    messages = [line["message"] for line in lines]
    assert any("Bắt đầu thực thi" in m for m in messages)
    assert any("Hoàn tất lần chạy" in m for m in messages)
    assert {line["level"] for line in lines} <= {"INFO", "WARN", "ERROR", "DEBUG"}
    assert payload.rstrip().endswith('data: {"run_id": ' + str(run_id) + "}")

    # Nối lại bằng Last-Event-ID: chỉ nhận các dòng sau id đó.
    last_id = lines[-2]["id"]
    with client.stream(
        "GET", f"/api/v1/runs/{run_id}/logs", headers={"Last-Event-ID": str(last_id)}
    ) as stream:
        resumed = "".join(stream.iter_text())
    assert resumed.count("event: log") == 1


def test_cancel_and_permissions(client: TestClient, user_client: TestClient, make_employee) -> None:
    employee = make_employee()
    assert user_client.post("/api/v1/runs", json={"employee_id": employee.id}).status_code == 403
    run_id = client.post("/api/v1/runs", json={"employee_id": employee.id}).json()["id"]
    assert user_client.post(f"/api/v1/runs/{run_id}/cancel").status_code == 403

    cancelled = client.post(f"/api/v1/runs/{run_id}/cancel")
    assert cancelled.status_code == 200 and cancelled.json()["status"] == "CANCELLED"
    assert client.post(f"/api/v1/runs/{run_id}/cancel").status_code == 409
    assert client.get("/api/v1/runs/99999").status_code == 404


def test_cannot_run_unapproved_or_archived_employee(client: TestClient, make_employee) -> None:
    draft = make_employee(approve=False)
    resp = client.post("/api/v1/runs", json={"employee_id": draft.id})
    assert resp.status_code == 409
    assert "chưa có quy trình được duyệt" in resp.json()["error"]["message"]
