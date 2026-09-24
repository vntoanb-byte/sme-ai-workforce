"""
Công cụ thị giác

  - vision.extract_invoice: đọc từng tệp hoá đơn bằng mô hình (dùng lại
    services.document_service.extract_document). Khử trùng theo sha256: tệp đã
    được xử lý xong ở lần trước thì bỏ qua; tệp dở dang (chưa có extraction)
    thì đọc lại — nhờ vậy thử lại cả lần chạy là an toàn.
  - vision.classify_document: phân loại tài liệu vào tập nhãn cho trước, xuất
    bảng phân loại Excel.

Bản ghi extraction luôn có model_name (từ phản hồi mô hình) và schema_version;
phiên bản prompt ghi trong nhật ký lần chạy.
"""

from __future__ import annotations

import io
import mimetypes
from typing import Any

from openpyxl import Workbook
from sqlalchemy import select

from app.models.artifact import Artifact, Document
from app.ports.llm import LLMInvalidOutput, LLMTimeout, LLMUnavailable
from app.services import document_service
from app.tools.base import (
    ToolBase,
    ToolContext,
    ToolResult,
    read_entry,
    register_tool,
    save_output,
    split_list,
)
from app.utils.hashing import sha256_bytes

_DONE_STATUSES = {"ok", "needs_review", "rejected"}
_FILE_ENTRY = {
    "type": "object",
    "properties": {
        "filename": {"type": "string"},
        "path": {"type": "string"},
        "artifact_id": {"type": "integer"},
    },
}


def _content_type(filename: str) -> str | None:
    guessed, _ = mimetypes.guess_type(filename)
    return guessed


@register_tool
class ExtractInvoice(ToolBase):
    code = "vision.extract_invoice"
    name = "Đọc hoá đơn bằng AI"
    category = "vision"
    description = "Trích xuất dữ liệu hoá đơn GTGT từ ảnh/PDF bằng mô hình thị giác."
    input_schema = {
        "type": "object",
        "properties": {"files": {"type": "array", "items": _FILE_ENTRY}},
        "required": ["files"],
    }
    output_schema = {
        "type": "object",
        "properties": {"document_ids": {"type": "array", "items": {"type": "integer"}}},
    }

    def run(self, ctx: ToolContext, config: dict[str, Any], inputs: dict[str, Any]) -> ToolResult:
        files: list[dict[str, Any]] = inputs.get("files") or []
        document_ids: list[int] = []
        duplicates = failed = 0
        for index, entry in enumerate(files, start=1):
            filename = str(entry.get("filename") or "tep")
            data = read_entry(ctx, entry)
            document = self._existing_document(ctx, data)
            if document is not None and any(
                e.run_id is not None and e.run_id == ctx.run_id for e in document.extractions
            ):
                # Đã đọc ở lần thử trước của CHÍNH lần chạy này — không gọi lại mô
                # hình nhưng vẫn chuyển tiếp cho bước QC (thử lại không bỏ sót).
                document_ids.append(document.id)
                ctx.log("INFO", f"[{index}/{len(files)}] {filename} · đã đọc ở lần thử trước")
                continue
            if document is not None and (
                document.status in _DONE_STATUSES or document.extractions
            ):
                duplicates += 1
                ctx.log("INFO", f"[{index}/{len(files)}] {filename} · đã xử lý trước đó, bỏ qua")
                continue

            kind = document_service.detect_source_kind(_content_type(filename))
            if kind is None:
                ctx.log("WARN", f"[{index}/{len(files)}] {filename} · không phải ảnh/PDF, bỏ qua")
                continue
            if document is None:
                document = Document(filename=filename, source_kind=kind, status="processing")
                document.artifact = document_service.ensure_artifact(
                    ctx.db, ctx.storage, data, filename
                )
                ctx.db.add(document)
                ctx.db.flush()
            # Commit TRƯỚC lời gọi mô hình (10–25 giây): SQLite chỉ có một người ghi —
            # giữ khoá ghi suốt lúc chờ mô hình sẽ làm API (tải lên, huỷ, xác nhận)
            # báo "database is locked" (lỗi thật phát hiện qua test huỷ giữa chừng).
            ctx.db.commit()

            extraction = document_service.extract_document(
                ctx.db,
                ctx.llm,
                document,
                data,
                run_id=ctx.run_id,
                run_qc=False,
                raise_transient=True,
            )
            ctx.db.flush()
            if extraction is None:
                failed += 1
                ctx.log(
                    "WARN",
                    f"[{index}/{len(files)}] {filename} · không đọc được ({document.status})",
                )
                continue
            document_ids.append(document.id)
            ctx.log(
                "INFO",
                f"[{index}/{len(files)}] HĐ {extraction.invoice_no or '?'} · đọc xong "
                f"{extraction.latency_ms / 1000:.1f}s",
            )

        detail = f"Đã đọc {len(document_ids)} hoá đơn"
        if duplicates:
            detail += f", bỏ qua {duplicates} tệp trùng"
        if failed:
            detail += f", {failed} tệp không đọc được"
        return ToolResult(outputs={"document_ids": document_ids}, detail=detail)

    @staticmethod
    def _existing_document(ctx: ToolContext, data: bytes) -> Document | None:
        return ctx.db.scalar(
            select(Document)
            .join(Artifact, Document.artifact_id == Artifact.id)
            .where(Artifact.sha256 == sha256_bytes(data))
            .order_by(Document.id.desc())
            .limit(1)
        )


@register_tool
class ClassifyDocument(ToolBase):
    code = "vision.classify_document"
    name = "Phân loại tài liệu bằng AI"
    category = "vision"
    description = "Phân loại từng tài liệu vào một trong các nhãn cho trước."
    input_schema = {
        "type": "object",
        "properties": {"files": {"type": "array", "items": _FILE_ENTRY}},
        "required": ["files"],
    }
    output_schema = {
        "type": "object",
        "properties": {
            "classifications": {"type": "array"},
            "artifact_ids": {"type": "array", "items": {"type": "integer"}},
        },
    }
    required_params = ("labels",)

    def run(self, ctx: ToolContext, config: dict[str, Any], inputs: dict[str, Any]) -> ToolResult:
        labels = split_list(config.get("labels"))
        if "khac" not in labels:
            labels.append("khac")
        schema = {
            "type": "object",
            "properties": {"label": {"type": "string", "enum": labels}},
            "required": ["label"],
            "additionalProperties": False,
        }
        prompt = (
            "Đây là ảnh một tài liệu. Hãy cho biết tài liệu thuộc loại nào trong các nhãn: "
            + ", ".join(labels)
            + ". Chỉ trả về JSON theo schema."
        )
        results: list[dict[str, str]] = []
        files: list[dict[str, Any]] = inputs.get("files") or []
        for entry in files:
            filename = str(entry.get("filename") or "tep")
            data = read_entry(ctx, entry)
            kind = document_service.detect_source_kind(_content_type(filename))
            label = "khac"
            if kind is not None:
                try:
                    image = document_service.prepare_image_bytes(data, kind)
                    result = ctx.llm.complete(
                        [{"role": "user", "content": prompt}], schema=schema, images=[image]
                    )
                    document_service.record_llm_call(ctx.db, result, run_id=ctx.run_id)
                    parsed = (result.parsed or {}).get("label")
                    label = parsed if parsed in labels else "khac"
                except (LLMTimeout, LLMUnavailable) as exc:
                    document_service.record_llm_call(ctx.db, None, run_id=ctx.run_id, error=exc)
                    raise
                except (LLMInvalidOutput, ValueError, OSError) as exc:
                    ctx.log("WARN", f"{filename} · không phân loại được: {exc}")
            results.append({"filename": filename, "label": label})
            ctx.log("INFO", f"{filename} → {label}")

        wb = Workbook()
        ws = wb.active
        ws.title = "PhanLoai"
        ws.append(["Tệp", "Loại tài liệu"])
        for row in results:
            ws.append([row["filename"], row["label"]])
        buf = io.BytesIO()
        wb.save(buf)
        artifact = save_output(ctx, buf.getvalue(), "BangPhanLoai.xlsx")
        counts: dict[str, int] = {}
        for row in results:
            counts[row["label"]] = counts.get(row["label"], 0) + 1
        summary = ", ".join(f"{k}: {v}" for k, v in sorted(counts.items())) or "không có tệp"
        return ToolResult(
            outputs={"classifications": results, "artifact_ids": [artifact.id]},
            detail=f"Đã phân loại {len(results)} tài liệu ({summary})",
        )
