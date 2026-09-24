"""
Công cụ bảng tính

xlsx.append_rows, xlsx.merge_files, xlsx.normalize, xlsx.dedupe, xlsx.reconcile.

Nguyên tắc:
  - KHÔNG BAO GIỜ ghi đè tệp gốc — kết quả luôn là tệp MỚI có dấu thời gian,
    lưu vào kho (artifact) và gắn vào lần chạy để người dùng tải về.
  - append_rows mở tệp đích bằng openpyxl rồi ghi bổ sung, nên giữ nguyên định
    dạng/công thức sẵn có của tệp đích trong bản mới.
  - Bảng tính được hiểu là trang tính đầu tiên (hoặc `sheet`), dòng 1 là tiêu đề.
"""

from __future__ import annotations

import io
import json
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePath
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from app.core.config import settings
from app.core.errors import ValidationFailed
from app.models.artifact import Artifact, Document
from app.models.extraction import Extraction
from app.tools.base import (
    ToolBase,
    ToolContext,
    ToolResult,
    artifact_ref,
    read_entry,
    register_tool,
    resolve_allowed_path,
    save_output,
    split_list,
)
from app.utils.dates import parse_vn_date
from app.utils.money import parse_money

_IDS = {"type": "array", "items": {"type": "integer"}}
_FILES = {"type": "array", "items": {"type": "object"}}
_WORKBOOK_OUT = {
    "type": "object",
    "properties": {"workbook": {"type": "object"}, "artifact_ids": _IDS},
}
INVOICE_HEADER = [
    "Số hoá đơn", "Ký hiệu", "Ngày lập", "Đơn vị bán", "MST bên bán",
    "Tiền hàng", "Thuế suất (%)", "Tiền thuế", "Tổng thanh toán", "Tệp gốc",
]


def _stamp() -> str:
    return datetime.now(ZoneInfo(settings.TIMEZONE)).strftime("%Y%m%d_%H%M%S")


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_table(data: bytes, sheet: str | None = None) -> tuple[list[str], list[list[Any]]]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
    if not rows:
        return [], []
    header = [str(c).strip() if c is not None else f"Cột {i + 1}" for i, c in enumerate(rows[0])]
    body = [(r + [None] * len(header))[: len(header)] for r in rows[1:]]
    return header, body


def _write_table(title: str, header: list[str], rows: list[list[Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for row in rows:
        ws.append(row)
    return wb


def _workbook_bytes(ctx: ToolContext, inputs: dict[str, Any]) -> bytes:
    """Bảng tính đầu vào: `workbook` (artifact từ bước trước) hoặc tệp đầu tiên."""
    workbook = inputs.get("workbook")
    if workbook:
        return read_entry(ctx, workbook)
    files = inputs.get("files") or []
    if not files:
        raise ValidationFailed("Không có bảng tính đầu vào.")
    return read_entry(ctx, files[0])


def _output(ctx: ToolContext, wb: Workbook, stem: str, inputs: dict[str, Any]) -> dict[str, Any]:
    filename = f"{stem}_{_stamp()}.xlsx"
    artifact = save_output(ctx, _to_bytes(wb), filename)
    return {
        "workbook": {"filename": filename, "artifact_id": artifact.id},
        "artifact_ids": [*inputs.get("artifact_ids", []), artifact.id],
    }


@register_tool
class AppendRows(ToolBase):
    code = "xlsx.append_rows"
    name = "Ghi hoá đơn vào bảng tính"
    category = "xlsx"
    description = "Ghi bổ sung dữ liệu hoá đơn vào bản sao mới của tệp Excel đích."
    input_schema = {
        "type": "object",
        "properties": {"document_ids": _IDS},
        "required": ["document_ids"],
    }
    output_schema = _WORKBOOK_OUT
    required_params = ("file",)

    def run(self, ctx: ToolContext, config: dict[str, Any], inputs: dict[str, Any]) -> ToolResult:
        target = str(config.get("file") or "SoHoaDon.xlsx")
        sheet_name = str(config.get("sheet") or "HoaDon")
        wb = self._open_target(target)
        ws: Worksheet = (
            wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
        )
        if ws.max_row == 1 and all(cell.value is None for cell in ws[1]):
            # Trang tính trống: ghi tiêu đề vào ĐÚNG dòng 1 (ws.append sau khi đã
            # chạm ô A1 sẽ ghi xuống dòng 2 — lỗi thật đã gặp).
            for col, title in enumerate(INVOICE_HEADER, start=1):
                ws.cell(1, col, title)
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1 and wb["Sheet"].max_row <= 1:
            del wb["Sheet"]

        written = 0
        for document_id in inputs.get("document_ids") or []:
            row = self._row(ctx, document_id)
            if row is None:
                continue
            ws.append(row)
            for col in (6, 8, 9):
                ws.cell(ws.max_row, col).number_format = "#,##0"
            written += 1

        out = _output(ctx, wb, PurePath(target).stem, inputs)
        ctx.log("INFO", f"Ghi {written} dòng vào {out['workbook']['filename']}")
        return ToolResult(
            outputs=out, detail=f"Ghi {written} hoá đơn vào {out['workbook']['filename']}"
        )

    @staticmethod
    def _open_target(target: str) -> Workbook:
        """Mở tệp đích nếu có (giữ định dạng); không có thì tạo bảng mới.
        Đường dẫn tương đối không có trên đĩa được hiểu là tên tệp đầu ra."""
        try:
            path = resolve_allowed_path(target)
        except ValidationFailed:
            if PurePath(target).is_absolute():
                raise
            return Workbook()
        if path.is_file():
            return load_workbook(path)
        return Workbook()

    @staticmethod
    def _row(ctx: ToolContext, document_id: int) -> list[Any] | None:
        document = ctx.db.get(Document, document_id)
        extraction = ctx.db.scalar(
            select(Extraction)
            .where(Extraction.document_id == document_id)
            .order_by(Extraction.created_at.desc(), Extraction.id.desc())
            .limit(1)
        )
        if document is None or extraction is None:
            return None
        try:
            form = json.loads(extraction.extracted_data_json).get("invoice_form")
        except (ValueError, TypeError):
            form = None
        return [
            extraction.invoice_no,
            form,
            extraction.issue_date,
            extraction.seller_name,
            extraction.seller_tax_code,
            extraction.subtotal,
            extraction.vat_rate,
            extraction.vat_amount,
            extraction.total,
            document.filename,
        ]


@register_tool
class MergeFiles(ToolBase):
    code = "xlsx.merge_files"
    name = "Gộp các tệp Excel"
    category = "xlsx"
    description = "Gộp trang tính đầu tiên của nhiều tệp Excel thành một bảng (theo tiêu đề cột)."
    input_schema = {"type": "object", "properties": {"files": _FILES}, "required": ["files"]}
    output_schema = _WORKBOOK_OUT

    def run(self, ctx: ToolContext, config: dict[str, Any], inputs: dict[str, Any]) -> ToolResult:
        header: list[str] = []
        rows: list[list[Any]] = []
        files = inputs.get("files") or []
        for entry in files:
            file_header, body = _read_table(read_entry(ctx, entry))
            for col in file_header:
                if col not in header:
                    header.append(col)
            index = {c: i for i, c in enumerate(file_header)}
            rows.extend([[r[index[c]] if c in index else None for c in header] for r in body])
        rows = [(r + [None] * len(header))[: len(header)] for r in rows]
        out = _output(ctx, _write_table("GopDuLieu", header, rows), "GopDuLieu", inputs)
        return ToolResult(outputs=out, detail=f"Gộp {len(files)} tệp → {len(rows)} dòng")


_SPACES = re.compile(r"\s+")
_DATE_LIKE = re.compile(r"^\d{1,2}[/.\-]\d{1,2}[/.\-]\d{4}$|^ng[àa]y\s", re.IGNORECASE)
_MONEY_LIKE = re.compile(r"^-?\d{1,3}([.,]\d{3})+([.,]\d+)?\s*(đ|vnđ|vnd|₫)?$", re.IGNORECASE)


def normalize_cell(value: Any) -> Any:
    """Bỏ khoảng trắng thừa; chuỗi dạng ngày → date; chuỗi dạng tiền → Decimal."""
    if not isinstance(value, str):
        return value
    text = _SPACES.sub(" ", value).strip()
    if not text:
        return None
    if _DATE_LIKE.search(text):
        parsed_date = parse_vn_date(text)
        if parsed_date is not None:
            return parsed_date
    if _MONEY_LIKE.match(text):
        amount = parse_money(text)
        if amount is not None:
            return amount
    return text


@register_tool
class Normalize(ToolBase):
    code = "xlsx.normalize"
    name = "Chuẩn hoá bảng tính"
    category = "xlsx"
    description = "Bỏ khoảng trắng thừa, chuẩn hoá ngày (dd/mm/yyyy) và số tiền (1.234.567)."
    input_schema = {"type": "object", "properties": {"workbook": {"type": "object"}}}
    output_schema = _WORKBOOK_OUT

    def run(self, ctx: ToolContext, config: dict[str, Any], inputs: dict[str, Any]) -> ToolResult:
        header, rows = _read_table(_workbook_bytes(ctx, inputs))
        changed = 0
        normalized: list[list[Any]] = []
        for row in rows:
            new_row = [normalize_cell(v) for v in row]
            changed += sum(1 for a, b in zip(row, new_row, strict=True) if a != b)
            normalized.append(new_row)
        wb = _write_table("ChuanHoa", header, normalized)
        ws = wb.active
        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                if isinstance(cell.value, date):
                    cell.number_format = "DD/MM/YYYY"
                elif isinstance(cell.value, Decimal):
                    cell.number_format = "#,##0.##"
        out = _output(ctx, wb, "ChuanHoa", inputs)
        return ToolResult(outputs=out, detail=f"Chuẩn hoá {changed} ô trên {len(rows)} dòng")


@register_tool
class Dedupe(ToolBase):
    code = "xlsx.dedupe"
    name = "Loại bỏ dòng trùng"
    category = "xlsx"
    description = "Bỏ các dòng trùng theo cột khoá (hoặc toàn dòng), giữ dòng xuất hiện đầu tiên."
    input_schema = {"type": "object", "properties": {"workbook": {"type": "object"}}}
    output_schema = _WORKBOOK_OUT

    def run(self, ctx: ToolContext, config: dict[str, Any], inputs: dict[str, Any]) -> ToolResult:
        header, rows = _read_table(_workbook_bytes(ctx, inputs))
        keys = split_list(config.get("key_columns"))
        missing = [k for k in keys if k not in header]
        if missing:
            raise ValidationFailed(f"Không có cột {missing} trong bảng tính.")
        indexes = [header.index(k) for k in keys] if keys else list(range(len(header)))
        seen: set[tuple[str, ...]] = set()
        kept: list[list[Any]] = []
        for row in rows:
            key = tuple(str(row[i]).strip().lower() if row[i] is not None else "" for i in indexes)
            if key in seen:
                continue
            seen.add(key)
            kept.append(row)
        removed = len(rows) - len(kept)
        out = _output(ctx, _write_table("KhongTrung", header, kept), "KhongTrung", inputs)
        return ToolResult(outputs=out, detail=f"Bỏ {removed} dòng trùng, còn {len(kept)} dòng")


@register_tool
class Reconcile(ToolBase):
    code = "xlsx.reconcile"
    name = "Đối chiếu hai bảng tính"
    category = "xlsx"
    description = "Ghép hai tệp theo cột khoá, liệt kê dòng chỉ có ở một bên và dòng lệch giá trị."
    input_schema = {"type": "object", "properties": {}}
    output_schema = {
        "type": "object",
        "properties": {"artifact_ids": _IDS, "mismatch_count": {"type": "integer"}},
    }
    required_params = ("left_file", "right_file", "key_column")

    def run(self, ctx: ToolContext, config: dict[str, Any], inputs: dict[str, Any]) -> ToolResult:
        key = str(config["key_column"]).strip()
        left_name = str(config["left_file"])
        right_name = str(config["right_file"])
        lh, lrows = _read_table(self._load(ctx, left_name))
        rh, rrows = _read_table(self._load(ctx, right_name))
        for header, name in ((lh, left_name), (rh, right_name)):
            if key not in header:
                raise ValidationFailed(f"Tệp '{name}' không có cột khoá '{key}'.")
        compare = split_list(config.get("compare_columns")) or [
            c for c in lh if c in rh and c != key
        ]
        left = {self._norm(r[lh.index(key)]): r for r in lrows}
        right = {self._norm(r[rh.index(key)]): r for r in rrows}

        only_left = [left[k] for k in left if k not in right]
        only_right = [right[k] for k in right if k not in left]
        diffs: list[list[Any]] = []
        for k in left.keys() & right.keys():
            for col in compare:
                if col not in lh or col not in rh:
                    continue
                a, b = left[k][lh.index(col)], right[k][rh.index(col)]
                if self._norm(a) != self._norm(b):
                    diffs.append([k, col, a, b])

        wb = Workbook()
        ws = wb.active
        ws.title = "LechGiaTri"
        ws.append([key, "Cột", f"Giá trị ({PurePath(left_name).name})",
                   f"Giá trị ({PurePath(right_name).name})"])
        for row in sorted(diffs, key=lambda r: (str(r[0]), str(r[1]))):
            ws.append(row)
        sides = (("ChiCoBenTrai", lh, only_left), ("ChiCoBenPhai", rh, only_right))
        for title, header, rows in sides:
            sheet = wb.create_sheet(title)
            sheet.append(header)
            for row in rows:
                sheet.append(row)
        filename = f"DoiChieu_{_stamp()}.xlsx"
        artifact = save_output(ctx, _to_bytes(wb), filename)
        mismatches = len(diffs) + len(only_left) + len(only_right)
        detail = (
            f"{len(diffs)} ô lệch, {len(only_left)} dòng chỉ có bên trái, "
            f"{len(only_right)} dòng chỉ có bên phải"
        )
        ctx.log("INFO" if mismatches == 0 else "WARN", f"Đối chiếu: {detail}")
        return ToolResult(
            outputs={"artifact_ids": [artifact.id], "mismatch_count": mismatches}, detail=detail
        )

    @staticmethod
    def _load(ctx: ToolContext, name: str) -> bytes:
        if name.startswith("artifact:"):
            artifact = ctx.db.get(Artifact, int(name.split(":", 1)[1]))
            if artifact is None:
                raise ValidationFailed(f"Không tìm thấy tệp {name}.")
            with ctx.storage.open(artifact_ref(artifact)) as fh:
                return fh.read()
        path = resolve_allowed_path(name)
        if not path.is_file():
            raise ValidationFailed(f"Không tìm thấy tệp '{name}'.")
        return path.read_bytes()

    @staticmethod
    def _norm(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            value = value.date()
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        normalized = normalize_cell(value) if isinstance(value, str) else value
        if isinstance(normalized, Decimal):
            return format(normalized.normalize(), "f")
        return str(normalized).strip().lower()
