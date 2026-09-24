"""
Dịch vụ báo cáo

Tổng hợp dữ liệu hoá đơn theo kỳ và kết xuất tệp Excel/PDF.

  - aggregate(): một câu SQL GROUP BY trên extraction MỚI NHẤT của từng chứng
    từ đã xác nhận (status=ok) — không nạp toàn bộ bản ghi vào bộ nhớ.
  - export(): sinh tệp, lưu vào kho, trả về (artifact, filename).

Chỉ tính chứng từ status=ok: bản đang chờ xác nhận chưa đáng tin, bản bị từ
chối không được đưa vào số liệu.

Lưu ý tiền tệ: SQLite không có kiểu thập phân thật — SUM trả số thực, được
làm tròn ngay về Decimal 2 chữ số (đủ chính xác với số tiền VND < 10^13).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal
from zoneinfo import ZoneInfo

import structlog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.errors import ValidationFailed
from app.models.artifact import Artifact, Document
from app.models.extraction import Extraction
from app.ports.storage import FileStorage
from app.services.document_service import ensure_artifact
from app.utils.money import format_money

logger = structlog.get_logger(__name__)

GroupBy = Literal["seller", "month", "vat_rate"]
GROUP_LABELS = {"seller": "Nhà cung cấp", "month": "Tháng", "vat_rate": "Thuế suất"}
_CENT = Decimal("0.01")


@dataclass(frozen=True)
class ReportRow:
    group: str
    doc_count: int
    subtotal: Decimal
    vat_amount: Decimal
    total: Decimal


@dataclass(frozen=True)
class ReportData:
    date_from: date
    date_to: date
    group_by: str
    rows: list[ReportRow]

    @property
    def grand_total(self) -> Decimal:
        return sum((r.total for r in self.rows), Decimal("0"))

    def to_dict(self) -> dict[str, Any]:
        return {
            "from": self.date_from.isoformat(),
            "to": self.date_to.isoformat(),
            "group_by": self.group_by,
            "rows": [
                {
                    "group": r.group,
                    "doc_count": r.doc_count,
                    "subtotal": r.subtotal,
                    "vat_amount": r.vat_amount,
                    "total": r.total,
                }
                for r in self.rows
            ],
            "grand_total": self.grand_total,
        }


def _money(value: Any) -> Decimal:
    if value is None:
        return Decimal("0.00")
    return Decimal(str(value)).quantize(_CENT)


def aggregate(db: Session, date_from: date, date_to: date, group_by: str) -> ReportData:
    if group_by not in GROUP_LABELS:
        raise ValidationFailed(f"Tiêu chí nhóm '{group_by}' không hợp lệ.")
    if date_from > date_to:
        raise ValidationFailed("Ngày bắt đầu phải trước hoặc bằng ngày kết thúc.")

    ranked = select(
        Extraction.document_id.label("document_id"),
        Extraction.issue_date.label("issue_date"),
        Extraction.seller_name.label("seller_name"),
        Extraction.vat_rate.label("vat_rate"),
        Extraction.subtotal.label("subtotal"),
        Extraction.vat_amount.label("vat_amount"),
        Extraction.total.label("total"),
        func.row_number()
        .over(
            partition_by=Extraction.document_id,
            order_by=(Extraction.created_at.desc(), Extraction.id.desc()),
        )
        .label("rn"),
    ).subquery()
    group_expr: Any
    if group_by == "seller":
        group_expr = func.coalesce(ranked.c.seller_name, "Không rõ")
    elif group_by == "month":
        group_expr = func.strftime("%Y-%m", ranked.c.issue_date)
    else:
        group_expr = ranked.c.vat_rate
    group_col = group_expr.label("grp")

    stmt = (
        select(
            group_col,
            func.count(ranked.c.document_id),
            func.sum(ranked.c.subtotal),
            func.sum(ranked.c.vat_amount),
            func.sum(ranked.c.total),
        )
        .select_from(ranked)
        .join(Document, and_(Document.id == ranked.c.document_id, ranked.c.rn == 1))
        .where(
            Document.status == "ok",
            ranked.c.issue_date >= date_from,
            ranked.c.issue_date <= date_to,
        )
        .group_by(group_col)
        .order_by(func.sum(ranked.c.total).desc())
    )
    rows = [
        ReportRow(
            group=_group_label(group_by, grp),
            doc_count=int(count),
            subtotal=_money(subtotal),
            vat_amount=_money(vat),
            total=_money(total),
        )
        for grp, count, subtotal, vat, total in db.execute(stmt)
    ]
    return ReportData(date_from, date_to, group_by, rows)


def _group_label(group_by: str, value: Any) -> str:
    if value is None:
        return "Không rõ"
    if group_by == "month":
        year, _, month = str(value).partition("-")
        return f"{month}/{year}"
    if group_by == "vat_rate":
        return f"{Decimal(str(value)).normalize():f}%"
    return str(value)


def resolve_period(period: str, today: date | None = None) -> tuple[date, date]:
    """'today' | 'this_week' | 'this_month' | 'last_month' → (từ ngày, đến ngày)
    theo múi giờ cấu hình (TIMEZONE)."""
    today = today or datetime.now(ZoneInfo(settings.TIMEZONE)).date()
    if period == "today":
        return today, today
    if period == "this_week":
        return today - timedelta(days=today.weekday()), today
    if period == "this_month":
        return today.replace(day=1), today
    if period == "last_month":
        last_day = today.replace(day=1) - timedelta(days=1)
        return last_day.replace(day=1), last_day
    raise ValidationFailed(f"Kỳ báo cáo '{period}' không hợp lệ.")


# ─── Kết xuất ───


def build_xlsx(data: ReportData) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "TongHop"
    ws.append([f"BÁO CÁO TỔNG HỢP HOÁ ĐƠN {data.date_from:%d/%m/%Y} – {data.date_to:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    header = [GROUP_LABELS[data.group_by], "Số chứng từ", "Tiền hàng", "Tiền thuế", "Tổng cộng"]
    ws.append(header)
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF7")
    for r in data.rows:
        ws.append([r.group, r.doc_count, r.subtotal, r.vat_amount, r.total])
    ws.append(
        [
            "Tổng cộng",
            sum(r.doc_count for r in data.rows),
            sum((r.subtotal for r in data.rows), Decimal("0")),
            sum((r.vat_amount for r in data.rows), Decimal("0")),
            data.grand_total,
        ]
    )
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "#,##0"
    for col, width in zip("ABCDE", (44, 14, 18, 16, 18), strict=True):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/Library/Fonts/Arial.ttf",
)
_font_name: str | None = None


def _pdf_font() -> str:
    """Đăng ký phông TTF có dấu tiếng Việt; không có thì dùng Helvetica (mất dấu)."""
    global _font_name
    if _font_name:
        return _font_name
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for candidate in (settings.REPORT_FONT_PATH, *_FONT_CANDIDATES):
        if candidate and Path(candidate).is_file():
            pdfmetrics.registerFont(TTFont("VNFont", candidate))
            _font_name = "VNFont"
            return _font_name
    logger.warning("report.font_missing", message="Không tìm thấy phông tiếng Việt cho PDF.")
    _font_name = "Helvetica"
    return _font_name


def build_pdf(data: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _pdf_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Báo cáo tổng hợp hoá đơn")
    title = ParagraphStyle("title", fontName=font, fontSize=14, leading=18)
    body = ParagraphStyle("body", fontName=font, fontSize=9, leading=12)
    table_rows: list[list[Any]] = [
        [GROUP_LABELS[data.group_by], "Số CT", "Tiền hàng", "Tiền thuế", "Tổng cộng"]
    ]
    for r in data.rows:
        table_rows.append(
            [
                Paragraph(r.group, body),
                str(r.doc_count),
                format_money(r.subtotal),
                format_money(r.vat_amount),
                format_money(r.total),
            ]
        )
    table_rows.append(
        [
            "Tổng cộng",
            str(sum(r.doc_count for r in data.rows)),
            format_money(sum((r.subtotal for r in data.rows), Decimal("0"))),
            format_money(sum((r.vat_amount for r in data.rows), Decimal("0"))),
            format_money(data.grand_total),
        ]
    )
    table = Table(table_rows, colWidths=[200, 45, 85, 75, 90], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F4F6FA")),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C6D2E6")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    doc.build(
        [
            Paragraph("BÁO CÁO TỔNG HỢP HOÁ ĐƠN", title),
            Paragraph(
                f"Kỳ: {data.date_from:%d/%m/%Y} – {data.date_to:%d/%m/%Y} · "
                f"Nhóm theo: {GROUP_LABELS[data.group_by].lower()}",
                body,
            ),
            Spacer(1, 10),
            table,
        ]
    )
    return buf.getvalue()


def export(
    db: Session, storage: FileStorage, data: ReportData, fmt: str
) -> tuple[Artifact, str]:
    """Sinh tệp báo cáo (xlsx|pdf), lưu kho; trả về (artifact, tên tệp gợi ý)."""
    if fmt == "xlsx":
        content = build_xlsx(data)
    elif fmt == "pdf":
        content = build_pdf(data)
    else:
        raise ValidationFailed(f"Định dạng '{fmt}' không được hỗ trợ (chỉ xlsx, pdf).")
    filename = f"BaoCao_{data.date_from:%Y%m%d}_{data.date_to:%Y%m%d}_{data.group_by}.{fmt}"
    return ensure_artifact(db, storage, content, filename), filename
